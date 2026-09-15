"""
Genera el manual de usuario de Polux en formato Word (.docx).

El manual se escribe con `python-docx` en vez de a mano para que pueda
regenerarse cuando la aplicación cambie: el texto vive en este script y
las capturas de pantalla en `docs/capturas/`, de modo que actualizar el
documento es volver a ejecutarlo y no rehacer la maquetación.

Uso:
    python scripts/generar_manual.py

Genera `docs/MANUAL_USUARIO.docx`. Requiere `python-docx`, que no forma
parte de las dependencias de la aplicación (solo se necesita para
producir documentación, no para ejecutar Polux).
"""

import os
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CAPTURAS = os.path.join(RAIZ, "docs", "capturas")
DESTINO = os.path.join(RAIZ, "docs", "MANUAL_USUARIO.docx")

#: Azul de la cabecera de las hojas de ruta en PDF, para que el manual y
#: lo que el usuario imprime se reconozcan como la misma herramienta.
AZUL_POLUX = RGBColor(0x1F, 0x3A, 0x68)

#: Ancho útil de la página con los márgenes definidos más abajo.
ANCHO_UTIL_CM = 16.0


# ---------------------------------------------------------------------------
# Utilidades de maquetación
# ---------------------------------------------------------------------------

def preparar_estilos(documento: Document) -> None:
    """Fija tipografía, márgenes y colores de los estilos que se usan."""
    normal = documento.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.15

    for nivel, tamano in ((1, 18), (2, 14), (3, 12)):
        estilo = documento.styles[f"Heading {nivel}"]
        estilo.font.name = "Calibri"
        estilo.font.size = Pt(tamano)
        estilo.font.color.rgb = AZUL_POLUX
        estilo.font.bold = True
        estilo.paragraph_format.space_before = Pt(16 if nivel == 1 else 12)
        estilo.paragraph_format.space_after = Pt(6)

    for seccion in documento.sections:
        seccion.top_margin = Cm(2.2)
        seccion.bottom_margin = Cm(2.2)
        seccion.left_margin = Cm(2.5)
        seccion.right_margin = Cm(2.5)


def parrafo(documento: Document, texto: str, **kwargs) -> None:
    """
    Añade un párrafo admitiendo **negrita** y `código` en el texto.

    Evita tener que trocear cada frase en runs a mano, que es lo que hace
    ilegible un generador de documentos.
    """
    p = documento.add_paragraph(**kwargs)
    for trozo in _trocear(texto):
        run = p.add_run(trozo["texto"])
        run.bold = trozo["negrita"]
        if trozo["codigo"]:
            run.font.name = "Consolas"
            run.font.size = Pt(10)
    return p


def _trocear(texto: str) -> list[dict]:
    """Parte el texto en fragmentos según los marcadores ** y `."""
    trozos: list[dict] = []
    actual = ""
    negrita = False
    codigo = False
    indice = 0
    while indice < len(texto):
        if texto.startswith("**", indice):
            if actual:
                trozos.append({"texto": actual, "negrita": negrita, "codigo": codigo})
                actual = ""
            negrita = not negrita
            indice += 2
        elif texto[indice] == "`":
            if actual:
                trozos.append({"texto": actual, "negrita": negrita, "codigo": codigo})
                actual = ""
            codigo = not codigo
            indice += 1
        else:
            actual += texto[indice]
            indice += 1
    if actual:
        trozos.append({"texto": actual, "negrita": negrita, "codigo": codigo})
    return trozos


def vineta(documento: Document, texto: str, nivel: int = 0) -> None:
    estilo = "List Bullet" if nivel == 0 else f"List Bullet {nivel + 1}"
    parrafo(documento, texto, style=estilo)


def lista_numerada(documento: Document, pasos: list[str]) -> None:
    """
    Escribe una lista numerada con el número puesto a mano.

    No se usa el estilo "List Number" de Word porque todas las listas del
    documento comparten entonces un mismo contador: la segunda lista
    empezaba en 3 en vez de en 1. Numerar a mano es menos elegante, pero
    cada lista empieza donde debe y se ve igual en cualquier procesador.
    """
    for numero, texto in enumerate(pasos, start=1):
        p = parrafo(documento, f"{numero}.\t{texto}")
        p.paragraph_format.left_indent = Cm(0.9)
        p.paragraph_format.first_line_indent = Cm(-0.9)
        p.paragraph_format.space_after = Pt(4)


def _recuadro(documento: Document, relleno: str) -> "object":
    """
    Crea una tabla de una sola celda, sin bordes y con fondo de color.

    Los recuadros no se hacen sombreando un párrafo: cuando uno de esos
    párrafos cae al final de una página, el color se estira hasta el
    margen inferior. Una celda de tabla mide exactamente lo que ocupa su
    contenido, y se comporta igual en Word, en Pages y en LibreOffice.
    """
    t = documento.add_table(rows=1, cols=1)
    t.autofit = False
    celda = t.rows[0].cells[0]
    celda.width = Cm(ANCHO_UTIL_CM)

    sombreado = OxmlElement("w:shd")
    sombreado.set(qn("w:val"), "clear")
    sombreado.set(qn("w:fill"), relleno)
    celda._tc.get_or_add_tcPr().append(sombreado)
    return celda


def _pegar_al_siguiente(documento: Document) -> None:
    """
    Impide que el último párrafo se quede solo al final de una página.

    Se usa antes de insertar un bloque de código o una figura, para que el
    texto que los presenta viaje con ellos y no acabe huérfano en la
    página anterior.
    """
    if documento.paragraphs:
        documento.paragraphs[-1].paragraph_format.keep_with_next = True


def bloque_codigo(documento: Document, lineas: list[str]) -> None:
    """Inserta un bloque monoespaciado con fondo gris."""
    _pegar_al_siguiente(documento)
    celda = _recuadro(documento, "F2F2F2")
    for indice, linea in enumerate(lineas):
        p = celda.paragraphs[0] if indice == 0 else celda.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(linea if linea else " ")
        run.font.name = "Consolas"
        run.font.size = Pt(9)
    documento.add_paragraph().paragraph_format.space_after = Pt(4)


def captura(documento: Document, archivo: str, pie: str, ancho_cm: float = ANCHO_UTIL_CM) -> None:
    """Inserta una captura centrada con su pie de figura."""
    ruta = os.path.join(CAPTURAS, archivo)
    if not os.path.isfile(ruta):
        raise SystemExit(f"Falta la captura {ruta}. Genera antes docs/capturas/.")

    _pegar_al_siguiente(documento)
    p = documento.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(ruta, width=Cm(ancho_cm))

    leyenda = documento.add_paragraph()
    leyenda.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = leyenda.add_run(pie)
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)


def tabla(documento: Document, cabeceras: list[str], filas: list[list[str]],
          anchos_cm: list[float] | None = None) -> None:
    """Inserta una tabla con la primera fila como cabecera en negrita."""
    t = documento.add_table(rows=1, cols=len(cabeceras))
    t.style = "Table Grid"
    t.autofit = False

    for celda, titulo in zip(t.rows[0].cells, cabeceras):
        celda.text = ""
        run = celda.paragraphs[0].add_run(titulo)
        run.bold = True
        run.font.size = Pt(10)
        sombreado = OxmlElement("w:shd")
        sombreado.set(qn("w:fill"), "1F3A68")
        celda._tc.get_or_add_tcPr().append(sombreado)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    for fila in filas:
        celdas = t.add_row().cells
        for celda, valor in zip(celdas, fila):
            celda.text = ""
            p = celda.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            for trozo in _trocear(valor):
                run = p.add_run(trozo["texto"])
                run.bold = trozo["negrita"]
                run.font.size = Pt(10)
                if trozo["codigo"]:
                    run.font.name = "Consolas"
                    run.font.size = Pt(9)

    if anchos_cm:
        for fila in t.rows:
            for celda, ancho in zip(fila.cells, anchos_cm):
                celda.width = Cm(ancho)

    documento.add_paragraph().paragraph_format.space_after = Pt(4)


def aviso(documento: Document, titulo: str, texto: str) -> None:
    """Recuadro destacado para advertencias y consejos."""
    celda = _recuadro(documento, "EEF3FA")
    p = celda.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    encabezado = p.add_run(f"{titulo}  ")
    encabezado.bold = True
    encabezado.font.color.rgb = AZUL_POLUX
    for trozo in _trocear(texto):
        run = p.add_run(trozo["texto"])
        run.bold = trozo["negrita"]
        if trozo["codigo"]:
            run.font.name = "Consolas"
            run.font.size = Pt(10)
    documento.add_paragraph().paragraph_format.space_after = Pt(4)


def pie_de_pagina(documento: Document) -> None:
    """Numera las páginas con un campo PAGE."""
    p = documento.sections[0].footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Polux 1.0.0 — Manual de usuario   ·   ")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(0x77, 0x77, 0x77)

    campo = p.add_run()
    campo.font.size = Pt(8)
    campo.font.color.rgb = RGBColor(0x77, 0x77, 0x77)
    inicio = OxmlElement("w:fldChar")
    inicio.set(qn("w:fldCharType"), "begin")
    instruccion = OxmlElement("w:instrText")
    instruccion.set(qn("xml:space"), "preserve")
    instruccion.text = "PAGE"
    fin = OxmlElement("w:fldChar")
    fin.set(qn("w:fldCharType"), "end")
    for elemento in (inicio, instruccion, fin):
        campo._r.append(elemento)


# ---------------------------------------------------------------------------
# Contenido del manual
# ---------------------------------------------------------------------------

def portada(doc: Document) -> None:
    for _ in range(4):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Polux")
    run.font.size = Pt(52)
    run.font.bold = True
    run.font.color.rgb = AZUL_POLUX

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Manual de usuario")
    run.font.size = Pt(22)
    run.font.color.rgb = RGBColor(0x44, 0x44, 0x44)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(
        "Optimización de rutas de reparto con ventanas horarias\n"
        "y cumplimiento del Reglamento (CE) nº 561/2006"
    )
    run.font.size = Pt(12)
    run.italic = True
    run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    for _ in range(2):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Versión 1.0.0")
    run.font.size = Pt(11)
    run.bold = True

    doc.add_page_break()


def indice(doc: Document) -> None:
    doc.add_heading("Contenido", level=1)
    secciones = [
        "1. Qué hace Polux",
        "2. Instalación y primer arranque",
        "3. El archivo de clientes (CSV)",
        "4. La ventana principal",
        "5. Configurar la jornada",
        "6. Cargar los clientes",
        "7. Optimizar las rutas",
        "8. Entender los resultados",
        "9. El mapa interactivo",
        "10. Hojas de ruta en PDF",
        "11. Enviar la ruta al conductor con Google Maps",
        "12. Exportar los resultados",
        "13. Problemas frecuentes",
        "14. Dónde guarda Polux sus archivos",
    ]
    for seccion in secciones:
        p = doc.add_paragraph(seccion)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.left_indent = Cm(0.5)
    doc.add_page_break()


def seccion_que_hace(doc: Document) -> None:
    doc.add_heading("1. Qué hace Polux", level=1)
    parrafo(
        doc,
        "Polux reparte un conjunto de clientes entre los vehículos de una flota y decide "
        "en qué orden visitarlos, buscando recorrer la menor distancia posible. Al hacerlo "
        "respeta tres cosas a la vez:",
    )
    vineta(doc, "**La capacidad de cada vehículo**: no se carga más mercancía de la que cabe.")
    vineta(
        doc,
        "**La ventana horaria de cada cliente**: si un cliente solo recibe de 09:00 a 14:00, "
        "no se planifica una llegada fuera de ese margen.",
    )
    vineta(
        doc,
        "**Los tiempos de conducción y descanso del Reglamento (CE) nº 561/2006**: pausa "
        "obligatoria de 45 minutos tras 4 h 30 de conducción continua, máximo de 9 horas "
        "diarias (ampliables a 10 h dos veces por semana) y tope de 56 horas semanales.",
    )
    parrafo(
        doc,
        "Las distancias y los tiempos **no son en línea recta**: se consultan por carretera, "
        "siguiendo las calles reales, igual que haría un navegador.",
    )
    parrafo(doc, "Con el plan ya calculado, Polux entrega:")
    vineta(doc, "Un mapa interactivo con cada ruta dibujada en un color.")
    vineta(doc, "Una hoja de ruta en PDF por vehículo, lista para imprimir y dar al conductor.")
    vineta(doc, "Un enlace de Google Maps por ruta, para abrirlo en el móvil.")
    vineta(doc, "Un archivo CSV con el resultado completo, para llevarlo a Excel.")

    aviso(
        doc,
        "Importante:",
        "Polux planifica **un viaje por vehículo y día**. Si la demanda total supera la "
        "capacidad de toda la flota, los clientes que no caben aparecen listados con su "
        "motivo, no se pierden en silencio.",
    )


def seccion_instalacion(doc: Document) -> None:
    doc.add_heading("2. Instalación y primer arranque", level=1)

    doc.add_heading("2.1. Windows", level=2)
    parrafo(doc, "No hace falta instalar Python ni ninguna otra cosa. Son dos pasos:")
    lista_numerada(doc, [
        "Descarga el archivo `Polux-1.0.0.exe` desde la página de versiones del proyecto.",
        "Haz doble clic. La aplicación arranca directamente; no hay instalador.",
    ])

    aviso(
        doc,
        "Aviso de Windows:",
        "La primera vez aparecerá la pantalla azul **«Windows ha protegido tu PC»**. Ocurre "
        "porque el programa no está firmado con un certificado digital de pago, no porque "
        "tenga nada dañino. Pulsa **«Más información»** y después **«Ejecutar de todas "
        "formas»**. El aviso no vuelve a salir.",
    )

    doc.add_heading("2.2. macOS y Linux (desde el código fuente)", level=2)
    parrafo(doc, "Con Python 3.10 o superior instalado:")
    bloque_codigo(doc, [
        "pip install -r requirements.txt",
        "python main.py",
    ])

    doc.add_heading("2.3. Conexión a internet", level=2)
    parrafo(doc, "Polux necesita red para dos cosas:")
    vineta(doc, "**Calcular distancias reales por carretera** entre el depósito y los clientes.")
    vineta(
        doc,
        "**Convertir direcciones de texto en coordenadas**, si tu archivo de clientes usa "
        "direcciones en vez de latitud y longitud.",
    )
    parrafo(
        doc,
        "Las direcciones ya consultadas alguna vez quedan guardadas, así que la segunda vez "
        "funcionan sin conexión. El cálculo de distancias, en cambio, siempre necesita red.",
    )


def seccion_csv(doc: Document) -> None:
    doc.add_heading("3. El archivo de clientes (CSV)", level=1)
    parrafo(
        doc,
        "Todo empieza aquí. Polux lee los clientes de un archivo **CSV**: una tabla de texto "
        "plano, separada por comas, que puedes preparar en Excel, en Google Sheets o en el "
        "Bloc de notas. La primera línea son siempre los nombres de las columnas.",
    )

    doc.add_heading("3.1. Columnas obligatorias", level=2)
    parrafo(doc, "Estas seis columnas tienen que estar siempre, escritas exactamente así:")
    tabla(
        doc,
        ["Columna", "Qué contiene", "Ejemplo"],
        [
            ["`id`", "Número entero que identifica al cliente. No se puede repetir, y **el 0 está reservado para el depósito**.", "`1`"],
            ["`nombre`", "Nombre del cliente. Es lo que aparece en el mapa y en la hoja de ruta.", "`Cliente Ruzafa`"],
            ["`demanda`", "Cuánto ocupa su pedido, en las mismas unidades que la capacidad del vehículo (kilos, palés, bultos… lo que decidas).", "`10`"],
            ["`hora_inicio`", "Primera hora a la que admite la entrega, en formato **HH:MM** de 24 horas.", "`08:00`"],
            ["`hora_fin`", "Última hora a la que admite la entrega. Tiene que ser posterior a `hora_inicio`.", "`12:00`"],
            ["`tiempo_servicio`", "Minutos que se tarda en descargar y entregar en ese cliente.", "`15`"],
        ],
        anchos_cm=[3.2, 9.6, 3.2],
    )

    doc.add_heading("3.2. La ubicación: dos formas de indicarla", level=2)
    parrafo(
        doc,
        "Además de esas seis, cada cliente necesita que se sepa dónde está. Tienes dos "
        "opciones y puedes elegir la que te resulte más cómoda:",
    )
    tabla(
        doc,
        ["Opción", "Columnas", "Cuándo usarla"],
        [
            ["**A. Coordenadas**", "`lat` y `lon`", "Cuando ya tienes las coordenadas. Es más rápido y no necesita consultar internet para localizar al cliente."],
            ["**B. Dirección**", "`direccion`", "Cuando solo tienes la dirección escrita. Polux la busca y la convierte en coordenadas automáticamente."],
        ],
        anchos_cm=[3.4, 3.0, 9.6],
    )
    parrafo(
        doc,
        "Si una fila trae **las dos cosas**, mandan las coordenadas: son inequívocas. El texto "
        "de la dirección no se desperdicia, se usa para que la hoja de ruta del conductor "
        "muestre una dirección legible en vez de un par de números.",
    )
    aviso(
        doc,
        "Consejo:",
        "Si vas a usar direcciones, escríbelas completas, con ciudad y país: "
        "`Carrer de Colón 25, València, España` se encuentra sin problema; `Colón 25`, no.",
    )

    doc.add_heading("3.3. Ejemplo con coordenadas", level=2)
    parrafo(doc, "Este es el archivo de ejemplo que viene con Polux, `clientes_ejemplo.csv`:")
    bloque_codigo(doc, [
        "id,nombre,lat,lon,demanda,hora_inicio,hora_fin,tiempo_servicio",
        "1,Cliente Ruzafa,39.4623,-0.3737,10,08:00,12:00,15",
        "2,Cliente Malvarrosa,39.4790,-0.3277,8,09:00,14:00,20",
        "3,Cliente Ciudad de las Artes,39.4540,-0.3540,15,08:30,13:00,10",
    ])

    doc.add_heading("3.4. Ejemplo con direcciones", level=2)
    bloque_codigo(doc, [
        "id,nombre,direccion,demanda,hora_inicio,hora_fin,tiempo_servicio",
        '1,Panadería Central,"Carrer de Colón 25, València, España",10,08:00,12:00,15',
        '2,Bar Malvarrosa,"Passeig Marítim 10, València, España",8,09:00,14:00,20',
    ])
    parrafo(
        doc,
        "Fíjate en las **comillas** alrededor de la dirección: son necesarias porque la "
        "dirección lleva comas dentro, y sin ellas el archivo se interpretaría mal.",
    )

    doc.add_heading("3.5. Reglas que conviene no olvidar", level=2)
    vineta(doc, "Los **decimales van con punto**, no con coma: `39.4623`, nunca `39,4623`.")
    vineta(doc, "Las **horas van en formato 24 h**: `14:30`, no `2:30 PM`.")
    vineta(doc, "Guarda el archivo como **CSV UTF-8** para que los acentos y las eñes salgan bien.")
    vineta(doc, "No dejes filas vacías al final de la tabla.")
    aviso(
        doc,
        "Desde Excel:",
        "Usa **Archivo → Guardar como → CSV UTF-8 (delimitado por comas)**. Si tu Excel está "
        "configurado en español, comprueba que los decimales se guardan con punto; si no, "
        "cámbialo en las opciones regionales antes de exportar.",
    )

    doc.add_heading("3.6. Qué pasa si algo está mal", level=2)
    parrafo(
        doc,
        "Polux revisa el archivo entero antes de empezar y, si encuentra un problema, dice "
        "**en qué línea** está. Estos son los mensajes más habituales:",
    )
    tabla(
        doc,
        ["Mensaje", "Qué significa"],
        [
            ["Faltan columnas obligatorias en el CSV", "Falta alguna de las seis columnas de la sección 3.1, o está mal escrita."],
            ["El CSV debe incluir 'lat' y 'lon', o bien una columna 'direccion'", "No hay ninguna forma de saber dónde están los clientes."],
            ["Línea N: el id X está duplicado", "Dos clientes comparten el mismo `id`."],
            ["Línea N: el id 0 está reservado para el depósito", "Ningún cliente puede llamarse `0`. Empieza a numerar en `1`."],
            ["Línea N: 'hora_inicio'/'hora_fin' deben tener formato HH:MM", "Alguna hora está escrita de otra forma."],
            ["Línea N: la ventana de tiempo es inválida", "`hora_fin` es anterior o igual a `hora_inicio`."],
            ["Línea N: coordenadas fuera de rango", "La latitud o la longitud no son valores geográficos válidos. Suele ser `lat` y `lon` intercambiadas."],
        ],
        anchos_cm=[7.0, 9.0],
    )
    parrafo(
        doc,
        "Hay un caso distinto: si una **dirección no se encuentra**, el archivo no se rechaza. "
        "El resto de clientes se cargan con normalidad y ese aparece después en la lista de "
        "clientes sin asignar, con el motivo `GEOCODIFICACIÓN`.",
    )
    doc.add_page_break()


def seccion_ventana(doc: Document) -> None:
    doc.add_heading("4. La ventana principal", level=1)
    parrafo(doc, "La pantalla se divide en tres zonas, y se usan de izquierda a derecha:")
    captura(doc, "01_ventana_inicial.png", "Figura 1. La ventana al arrancar, antes de cargar ningún cliente.")
    tabla(
        doc,
        ["Zona", "Para qué sirve"],
        [
            ["**Izquierda — Configuración**", "Dónde se dice cómo es la jornada: el depósito, la flota, los horarios. Debajo, un resumen que se actualiza solo."],
            ["**Centro — Mapa**", "El botón que abre el mapa interactivo en el navegador."],
            ["**Derecha — Resultados**", "Vacío hasta que optimizas. Después muestra las rutas y todos los botones de exportación."],
        ],
        anchos_cm=[5.0, 11.0],
    )
    aviso(
        doc,
        "Si no ves algún botón:",
        "Las columnas laterales **se desplazan**. Si la ventana es baja, aparece una barra de "
        "desplazamiento en el lateral de la columna; también puedes usar la rueda del ratón "
        "situando el puntero encima de la columna.",
    )


def seccion_configurar(doc: Document) -> None:
    doc.add_heading("5. Configurar la jornada", level=1)
    captura(doc, "02_panel_configuracion.png",
            "Figura 2. El panel de configuración, con los valores por defecto.", 10.5)

    tabla(
        doc,
        ["Campo", "Qué significa"],
        [
            ["**Dirección del depósito**", "El punto de donde salen todos los vehículos. Escribe la dirección y pulsa **Buscar**: debajo aparecerán las coordenadas encontradas. Viene rellenado con una dirección de ejemplo en València, así que la aplicación funciona nada más abrirla."],
            ["**Número de vehículos**", "Cuántos vehículos hay disponibles hoy."],
            ["**Flota ilimitada**", "Casilla para evaluación: permite usar tantos vehículos como haga falta. Sirve para saber cuántos necesitarías idealmente. Al marcarla, el número de vehículos se desactiva."],
            ["**Capacidad por vehículo**", "Cuánto carga un vehículo, en las mismas unidades que la columna `demanda` del CSV. Todos los vehículos tienen la misma."],
            ["**Hora de salida**", "A qué hora salen del depósito. Por defecto, 08:00."],
            ["**Hora límite de regreso**", "La hora a la que la jornada debe haber terminado. Ninguna ruta se planifica más allá. Por defecto, 18:00."],
            ["**Horas ya conducidas esta semana**", "Lo que el conductor lleva acumulado. Polux lo descuenta del tope semanal de 56 horas del Reglamento."],
            ["**Permitir jornada extendida**", "Sube el máximo diario de 9 a 10 horas. El Reglamento solo lo permite **dos veces por semana**."],
            ["**Contar el regreso al depósito**", "Marcada, las rutas terminan volviendo al depósito y ese trayecto cuenta para el horario. Desmarcada, la ruta acaba en el último cliente."],
            ["**Nombre del conductor**", "Opcional. Solo se usa para que aparezca impreso en la hoja de ruta."],
        ],
        anchos_cm=[5.2, 10.8],
    )
    aviso(
        doc,
        "Ojo:",
        "La hora límite de regreso tiene que ser **posterior** a la de salida. Si no, Polux lo "
        "avisa y no deja continuar.",
    )


def seccion_cargar(doc: Document) -> None:
    doc.add_heading("6. Cargar los clientes", level=1)
    parrafo(
        doc,
        "Pulsa **«Cargar clientes (CSV)…»** y elige tu archivo. Debajo del botón aparecerá "
        "el nombre del archivo y cuántos clientes se han leído.",
    )
    parrafo(
        doc,
        "Si el archivo usa direcciones, se abrirá una ventana de progreso mientras se "
        "localizan. Tarda **aproximadamente un segundo por dirección nueva**: es el ritmo "
        "máximo que permite el servicio gratuito de búsqueda que usa Polux. Las direcciones "
        "ya conocidas son instantáneas.",
    )
    parrafo(doc, "En cuanto se cargan, el resumen de la parte inferior izquierda se rellena:")
    captura(doc, "04_resumen.png", "Figura 3. El resumen previo, ya con los clientes cargados.", 10.5)
    tabla(
        doc,
        ["Dato", "Qué te está diciendo"],
        [
            ["Horas conducidas esta semana", "Lo que has escrito arriba, para tenerlo a la vista."],
            ["Techo diario disponible", "Cuánto se puede conducir hoy sin incumplir el Reglamento."],
            ["Horario de la jornada", "La franja entre la salida y el límite de regreso."],
            ["Clientes cargados", "Cuántos se han leído del archivo."],
            ["Factibilidad estimada", "Un cálculo rápido y aproximado de a cuántos clientes se puede llegar a tiempo. **Es una estimación**, no el resultado: el número definitivo sale al optimizar."],
            ["Advertencias", "Avisos previos, por ejemplo si la demanda total ya supera la capacidad de la flota."],
        ],
        anchos_cm=[5.2, 10.8],
    )


def seccion_optimizar(doc: Document) -> None:
    doc.add_heading("7. Optimizar las rutas", level=1)
    parrafo(
        doc,
        "Pulsa **«Optimizar rutas»**. El puntero cambia mientras Polux consulta las "
        "distancias por carretera y calcula el plan; con 15 clientes son unos segundos.",
    )
    parrafo(doc, "El cálculo tiene tres fases, que ocurren solas:")
    lista_numerada(doc, [
        "**Construcción**: se agrupan clientes en rutas fusionando las que más kilómetros ahorran.",
        "**Mejora**: se reordenan las paradas dentro de cada ruta para acortarla.",
        "**Reinserción**: se intenta colocar a los clientes que se habían quedado fuera en algún "
        "hueco todavía libre.",
    ])
    aviso(
        doc,
        "Si falla la conexión:",
        "Si el servicio de distancias no responde, Polux lo dice con un mensaje claro y **no "
        "calcula nada**. Prefiere no darte un plan antes que darte uno basado en distancias "
        "en línea recta, que no se parecen a la realidad.",
    )


def seccion_resultados(doc: Document) -> None:
    doc.add_heading("8. Entender los resultados", level=1)
    captura(doc, "09_ruta_seleccionada.png",
            "Figura 4. El panel de resultados con la primera ruta seleccionada.", 11.5)

    doc.add_heading("8.1. Los totales", level=2)
    vineta(doc, "**Distancia total**: kilómetros por carretera que recorre toda la flota.")
    vineta(doc, "**Vehículos utilizados**: cuántos hacen falta de verdad. Puede ser menos de los que tienes.")
    vineta(doc, "**Tiempo de espera total**: horas que los conductores pasan parados esperando a que abra la ventana de un cliente. Es la señal más útil de que las ventanas horarias están apretadas.")
    vineta(doc, "**Reinserción**: cuántos clientes recuperó la tercera fase.")

    doc.add_heading("8.2. La tabla de rutas", level=2)
    parrafo(doc, "Una fila por vehículo. **Haz clic en una fila** y debajo verás su secuencia completa: el orden de las paradas, la hora estimada de llegada a cada una, su ventana horaria y su demanda.")
    parrafo(doc, "Presta atención a la diferencia entre dos columnas:")
    vineta(doc, "**Tiempo de conducción** es lo que cuenta para el Reglamento.")
    vineta(doc, "**Tiempo total** incluye además las descargas y las esperas. Es lo que tarda la jornada de verdad.")

    doc.add_heading("8.3. Clientes sin asignar", level=2)
    parrafo(
        doc,
        "Si algún cliente se queda fuera, aparece aquí con el motivo. Nunca desaparece sin "
        "explicación:",
    )
    tabla(
        doc,
        ["Motivo", "Qué ha pasado", "Qué puedes hacer"],
        [
            ["`CAPACIDAD`", "No cabe en ningún vehículo con lo que ya llevan.", "Añadir un vehículo o subir la capacidad."],
            ["`VENTANA`", "No se llega a tiempo a su franja horaria.", "Ampliar su ventana en el CSV, o adelantar la hora de salida."],
            ["`TIEMPO`", "Atenderlo pasaría del límite diario de conducción o de la hora de regreso.", "Marcar la jornada extendida, o repartirlo en otro día."],
            ["`FLOTA`", "No quedan vehículos libres.", "Aumentar el número de vehículos."],
            ["`GEOCODIFICACIÓN`", "No se encontró su dirección.", "Revisarla y completarla con ciudad y país, o poner `lat` y `lon`."],
        ],
        anchos_cm=[3.6, 6.6, 5.8],
    )

    doc.add_heading("8.4. Advertencias del Reglamento", level=2)
    parrafo(
        doc,
        "El recuadro amarillo avisa cuando una ruta queda cerca de algún límite legal, o "
        "cuando lleva pausas obligatorias de 45 minutos. Una ruta con advertencia es legal; "
        "simplemente tiene poco margen.",
    )
    doc.add_page_break()


def seccion_mapa(doc: Document) -> None:
    doc.add_heading("9. El mapa interactivo", level=1)
    parrafo(doc, "Pulsa **«Abrir mapa interactivo en el navegador»**, en la zona central.")
    captura(doc, "08_mapa_navegador.png", "Figura 5. El mapa de rutas abierto en el navegador.")
    vineta(doc, "La casa negra es el **depósito**.")
    vineta(doc, "Cada **ruta** tiene su color y sigue las calles reales, no líneas rectas.")
    vineta(doc, "Cada **parada** está numerada. Pulsa un marcador para ver el cliente, su hora estimada de llegada, su ventana y su demanda.")
    vineta(doc, "Los **clientes sin asignar** salen en rojo, con un triángulo de aviso.")
    aviso(
        doc,
        "¿Por qué se abre fuera?",
        "El mapa se abre en tu navegador y no dentro de la ventana de Polux porque necesita "
        "JavaScript, y el visor que Polux lleva integrado no lo ejecuta. Abrirlo fuera es la "
        "forma de que funcione completo: con zoom, con desplazamiento y con los datos de cada "
        "marcador.",
    )


def seccion_pdf(doc: Document) -> None:
    doc.add_heading("10. Hojas de ruta en PDF", level=1)
    parrafo(
        doc,
        "Es lo que se le da impreso al conductor. Pulsa **«Hojas de ruta (PDF)…»** y se abre "
        "una vista previa con una hoja por vehículo:",
    )
    captura(doc, "07_dialogo_pdf.png", "Figura 6. Vista previa de la hoja de ruta.", 11.0)
    parrafo(doc, "Cada hoja lleva:")
    vineta(doc, "La **cabecera**: fecha, número de ruta, vehículo y capacidad, conductor, hora de salida, límite de regreso y dirección del depósito.")
    vineta(doc, "El **resumen**: número de paradas, distancia, hora de salida, tiempo de conducción, tiempo de espera y tiempo total estimado.")
    vineta(doc, "El **itinerario**: cada parada con su hora estimada de llegada, su demanda y el tiempo de servicio, **incluidas las pausas obligatorias** cuando la ruta las necesita, y la vuelta al depósito si está activada.")
    vineta(doc, "La lista de **clientes sin asignar** con sus motivos, para que quede constancia.")
    vineta(doc, "Una columna **«Observaciones» en blanco a propósito**, para que el conductor anote a mano.")
    parrafo(doc, "Con los botones de abajo:")
    vineta(doc, "**Siguiente / Anterior** recorren todas las páginas de todas las rutas.")
    vineta(doc, "**Guardar PDF** pide una carpeta y escribe un archivo por ruta.")
    vineta(doc, "**Imprimir** manda las hojas a la impresora predeterminada.")
    aviso(
        doc,
        "Direcciones en la hoja:",
        "En la columna «Dirección» aparece el texto de la dirección si tu CSV traía la columna "
        "`direccion`. Si solo diste `lat` y `lon`, se imprimen las coordenadas, que a un "
        "conductor le dicen poco. Si vas a imprimir hojas de ruta, merece la pena incluir "
        "también la columna `direccion`.",
    )


def seccion_google_maps(doc: Document) -> None:
    doc.add_heading("11. Enviar la ruta al conductor con Google Maps", level=1)
    parrafo(
        doc,
        "Hay un botón **«Abrir Ruta N en Google Maps»** por cada ruta. Al pulsarlo pasan dos "
        "cosas a la vez: el enlace se abre en el navegador y **se copia al portapapeles**.",
    )
    parrafo(
        doc,
        "Eso segundo es lo práctico: pégalo en un WhatsApp al conductor y, al abrirlo en el "
        "móvil, tendrá la navegación con todas las paradas en orden.",
    )


def seccion_exportar(doc: Document) -> None:
    doc.add_heading("12. Exportar los resultados", level=1)
    tabla(
        doc,
        ["Botón", "Qué genera"],
        [
            ["**Exportar resultados (CSV)**", "Un archivo para abrir en Excel, con una fila por parada y una fila por cliente sin asignar. La columna `tipo` distingue unas de otras y la columna `motivo` explica cada exclusión."],
            ["**Exportar mapa (HTML)**", "El mapa como archivo independiente. Se puede enviar por correo y se abre en cualquier navegador, sin necesidad de tener Polux."],
        ],
        anchos_cm=[5.2, 10.8],
    )


def seccion_problemas(doc: Document) -> None:
    doc.add_heading("13. Problemas frecuentes", level=1)
    tabla(
        doc,
        ["Situación", "Causa y solución"],
        [
            ["Windows avisa de que el programa no es seguro", "El ejecutable no está firmado con un certificado de pago. Pulsa «Más información» → «Ejecutar de todas formas»."],
            ["«No se pudo obtener la matriz de distancias»", "Sin conexión, o el servicio público de rutas no responde en ese momento. Comprueba la red y vuelve a intentarlo en unos minutos."],
            ["La carga del CSV va muy lenta", "Se están buscando direcciones, a un segundo cada una. Es el límite del servicio gratuito. La próxima vez irá rápido porque quedan guardadas."],
            ["No encuentra una dirección", "Escríbela más completa, con ciudad y país. Si sigue sin salir, pon `lat` y `lon` directamente."],
            ["Muchos clientes con motivo `TIEMPO`", "La jornada se queda corta. Prueba a adelantar la hora de salida, retrasar el límite de regreso o marcar la jornada extendida."],
            ["Muchos clientes con motivo `FLOTA`", "Faltan vehículos. Sube el número, o marca «Flota ilimitada» para ver cuántos harían falta."],
            ["No veo algunos botones", "La ventana es baja: desplaza la columna con la rueda del ratón o con su barra lateral."],
            ["El mapa sale en blanco en la ventana", "Es lo normal: el mapa se ve en el navegador, con el botón «Abrir mapa interactivo en el navegador»."],
        ],
        anchos_cm=[5.4, 10.6],
    )


def seccion_archivos(doc: Document) -> None:
    doc.add_heading("14. Dónde guarda Polux sus archivos", level=1)
    parrafo(doc, "La aplicación no ensucia la carpeta desde la que se ejecuta. Sus datos van a:")
    tabla(
        doc,
        ["Sistema", "Carpeta"],
        [
            ["Windows", "`C:\\Users\\<tu usuario>\\AppData\\Local\\Polux`"],
            ["macOS y Linux", "`~/.polux`"],
        ],
        anchos_cm=[4.0, 12.0],
    )
    parrafo(
        doc,
        "Ahí se guarda la memoria de direcciones ya buscadas. Los PDF y los CSV que exportes "
        "van a donde tú elijas en cada momento.",
    )


def main() -> int:
    if not os.path.isdir(CAPTURAS):
        print(f"No existe {CAPTURAS}: genera antes las capturas de pantalla.", file=sys.stderr)
        return 1

    doc = Document()
    preparar_estilos(doc)
    pie_de_pagina(doc)

    portada(doc)
    indice(doc)
    seccion_que_hace(doc)
    seccion_instalacion(doc)
    seccion_csv(doc)
    seccion_ventana(doc)
    seccion_configurar(doc)
    seccion_cargar(doc)
    seccion_optimizar(doc)
    seccion_resultados(doc)
    seccion_mapa(doc)
    seccion_pdf(doc)
    seccion_google_maps(doc)
    seccion_exportar(doc)
    seccion_problemas(doc)
    seccion_archivos(doc)

    doc.save(DESTINO)
    print(f"Manual generado: {DESTINO}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
