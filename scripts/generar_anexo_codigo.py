"""
Genera el anexo A de la memoria (fragmentos de código) en formato Word.

El anexo no se copia y pega a mano: los fragmentos se extraen del código
fuente en el momento de generar el documento, indicando el archivo y las
líneas originales. Así el anexo no puede quedar desfasado respecto al
código sin que se note, y regenerarlo es volver a ejecutar este script.

Uso:
    python scripts/generar_anexo_codigo.py

Genera `docs/ANEXO_A_CODIGO.docx`. Requiere `python-docx`, que solo se
necesita para producir documentación, no para ejecutar Polux.
"""

import os
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DESTINO = os.path.join(RAIZ, "docs", "ANEXO_A_CODIGO.docx")

#: Azul de la cabecera de las hojas de ruta, para que toda la documentación
#: del proyecto comparta el mismo color.
AZUL_POLUX = RGBColor(0x1F, 0x3A, 0x68)

#: Gris de los números de línea: presentes pero sin competir con el código.
GRIS_NUMERO = RGBColor(0x90, 0x90, 0x90)

#: Cuerpo del código. Con márgenes de 2 cm y la sangría de los números de
#: línea, a 9 pt caben unos 89 caracteres por línea en Consolas, y unos 82
#: si el equipo no tiene Consolas y el procesador la sustituye por una
#: monoespaciada más ancha como Courier. Las pocas líneas del proyecto que
#: superan ese ancho se parten con sangría francesa, de modo que la
#: continuación queda alineada con la columna del código y sin número, que
#: es lo que la distingue de una línea nueva.
TAMANO_CODIGO_PT = 9


# ---------------------------------------------------------------------------
# Extracción de los fragmentos desde el código fuente
# ---------------------------------------------------------------------------

def extraer(ruta_relativa: str, primera: int, ultima: int) -> list[str]:
    """
    Devuelve las líneas [primera, ultima] de un archivo, ambas incluidas.

    Los números son los del archivo original (base 1), tal y como los
    muestra un editor, para que el anexo pueda citarlos con exactitud.
    """
    ruta = os.path.join(RAIZ, ruta_relativa)
    with open(ruta, encoding="utf-8") as archivo:
        lineas = archivo.read().splitlines()

    if ultima > len(lineas):
        raise SystemExit(
            f"{ruta_relativa} tiene {len(lineas)} líneas; se pidió hasta la {ultima}."
        )
    return lineas[primera - 1 : ultima]


# ---------------------------------------------------------------------------
# Utilidades de maquetación
# ---------------------------------------------------------------------------

def corregir_ajustes(documento: Document) -> None:
    """
    Completa el elemento `w:zoom` que la plantilla de python-docx deja
    incompleto.

    La plantilla por defecto escribe `<w:zoom>` sin el atributo
    `w:percent`, que el esquema de OOXML exige. Los procesadores de texto
    lo toleran, pero un validador contra el esquema lo marca como error, y
    un anexo de una memoria académica conviene que valide limpio.
    """
    ajustes = documento.settings.element
    zoom = ajustes.find(qn("w:zoom"))
    if zoom is not None and zoom.get(qn("w:percent")) is None:
        zoom.set(qn("w:percent"), "100")


def preparar_estilos(documento: Document) -> None:
    """Fija tipografía, márgenes y colores."""
    normal = documento.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.15

    for nivel, tamano in ((1, 18), (2, 14)):
        estilo = documento.styles[f"Heading {nivel}"]
        estilo.font.name = "Calibri"
        estilo.font.size = Pt(tamano)
        estilo.font.color.rgb = AZUL_POLUX
        estilo.font.bold = True
        estilo.paragraph_format.space_before = Pt(18 if nivel == 1 else 14)
        estilo.paragraph_format.space_after = Pt(6)

    for seccion in documento.sections:
        seccion.top_margin = Cm(2.2)
        seccion.bottom_margin = Cm(2.2)
        seccion.left_margin = Cm(2.0)
        seccion.right_margin = Cm(2.0)


def _trocear(texto: str) -> list[dict]:
    """Parte el texto en fragmentos según los marcadores ** y `."""
    trozos: list[dict] = []
    actual = ""
    negrita = False
    codigo = False
    indice = 0
    while indice < len(texto):
        if texto.startswith("**", indice):
            if actual:
                trozos.append({"texto": actual, "negrita": negrita, "codigo": codigo})
                actual = ""
            negrita = not negrita
            indice += 2
        elif texto[indice] == "`":
            if actual:
                trozos.append({"texto": actual, "negrita": negrita, "codigo": codigo})
                actual = ""
            codigo = not codigo
            indice += 1
        else:
            actual += texto[indice]
            indice += 1
    if actual:
        trozos.append({"texto": actual, "negrita": negrita, "codigo": codigo})
    return trozos


def parrafo(documento: Document, texto: str, **kwargs):
    """Añade un párrafo admitiendo **negrita** y `código` en el texto."""
    p = documento.add_paragraph(**kwargs)
    for trozo in _trocear(texto):
        run = p.add_run(trozo["texto"])
        run.bold = trozo["negrita"]
        if trozo["codigo"]:
            run.font.name = "Consolas"
            run.font.size = Pt(10)
    return p


def procedencia(documento: Document, ruta: str, primera: int, ultima: int,
                nota: str = "") -> None:
    """
    Escribe la línea que identifica el origen exacto del fragmento.

    Es lo que permite al lector del anexo ir al archivo y encontrar el
    código, y lo que deja claro que el fragmento es una cita literal y no
    una reescritura para la memoria.
    """
    p = documento.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.space_before = Pt(6)

    etiqueta = p.add_run("Archivo: ")
    etiqueta.bold = True
    etiqueta.font.size = Pt(9)

    origen = p.add_run(f"{ruta}")
    origen.font.name = "Consolas"
    origen.font.size = Pt(9)

    rango = p.add_run(f"  ·  líneas {primera}–{ultima} del original")
    rango.font.size = Pt(9)
    rango.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    if nota:
        aclaracion = p.add_run(f"  ·  {nota}")
        aclaracion.font.size = Pt(9)
        aclaracion.italic = True
        aclaracion.font.color.rgb = RGBColor(0x55, 0x55, 0x55)


def _filete_izquierdo(parrafo_codigo) -> None:
    """
    Dibuja una línea vertical gris en el margen izquierdo del párrafo.

    Aplicada a todas las líneas de un bloque, las une visualmente en un
    único filete continuo que delimita el código sin necesidad de fondo.

    Se delimitan así los bloques, y no metiéndolos en una tabla de una
    celda, porque una fila de tabla no se reparte entre dos páginas en
    todos los procesadores: un fragmento más largo que el hueco restante
    se veía cortado, perdiendo las últimas líneas sin previo aviso.
    """
    bordes = OxmlElement("w:pBdr")
    izquierdo = OxmlElement("w:left")
    izquierdo.set(qn("w:val"), "single")
    izquierdo.set(qn("w:sz"), "12")      # 1,5 pt (la unidad es 1/8 de punto)
    izquierdo.set(qn("w:space"), "8")    # separación con el texto, en puntos
    izquierdo.set(qn("w:color"), "B0B8C4")
    bordes.append(izquierdo)

    # El orden de los hijos de <w:pPr> lo impone el esquema de OOXML, y
    # <w:pBdr> va antes que <w:spacing> y <w:ind>. Añadirlo al final deja
    # un documento que los procesadores no rechazan pero cuyo borde
    # simplemente ignoran, así que se inserta en su sitio.
    propiedades = parrafo_codigo._p.get_or_add_pPr()
    estilo = propiedades.find(qn("w:pStyle"))
    posicion = list(propiedades).index(estilo) + 1 if estilo is not None else 0
    propiedades.insert(posicion, bordes)


def _texto_literal(run, texto: str) -> None:
    """
    Escribe texto conservando la sangría exacta del código.

    Sin `xml:space="preserve"` los procesadores de texto pueden colapsar
    los espacios iniciales, y en Python la sangría *es* la sintaxis.
    """
    run.text = texto
    for elemento in run._r.findall(qn("w:t")):
        elemento.set(qn("xml:space"), "preserve")


def bloque_codigo(documento: Document, lineas: list[str]) -> None:
    """
    Inserta un bloque de código numerado desde 1.

    La numeración es relativa al fragmento, no al archivo: quien lee el
    anexo necesita poder referirse a "la línea 12 del fragmento A.2" sin
    tener delante el archivo completo. El número de línea del original
    queda indicado aparte, en la línea de procedencia.
    """
    ancho_numero = len(str(len(lineas)))
    #: Hueco reservado para el número de línea, para que el código de todas
    #: las líneas empiece a la misma altura.
    sangria_codigo = Cm(0.6 + 0.19 * (ancho_numero + 2))

    for indice, linea in enumerate(lineas, start=1):
        p = documento.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        # Sangría francesa: si una línea muy larga no cabe en el ancho de
        # la página, su continuación queda indentada y se distingue a
        # simple vista de una línea nueva del código.
        p.paragraph_format.left_indent = sangria_codigo
        p.paragraph_format.first_line_indent = -sangria_codigo + Cm(0.6)
        _filete_izquierdo(p)

        numero = p.add_run(str(indice).rjust(ancho_numero) + "  ")
        numero.font.name = "Consolas"
        numero.font.size = Pt(TAMANO_CODIGO_PT)
        numero.font.color.rgb = GRIS_NUMERO

        codigo = p.add_run()
        codigo.font.name = "Consolas"
        codigo.font.size = Pt(TAMANO_CODIGO_PT)
        _texto_literal(codigo, linea if linea else " ")

    documento.add_paragraph().paragraph_format.space_after = Pt(4)


def pie_de_pagina(documento: Document) -> None:
    """Numera las páginas con un campo PAGE."""
    p = documento.sections[0].footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Anexo A — Fragmentos de código de Polux   ·   ")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(0x77, 0x77, 0x77)

    campo = p.add_run()
    campo.font.size = Pt(8)
    campo.font.color.rgb = RGBColor(0x77, 0x77, 0x77)
    inicio = OxmlElement("w:fldChar")
    inicio.set(qn("w:fldCharType"), "begin")
    instruccion = OxmlElement("w:instrText")
    instruccion.set(qn("xml:space"), "preserve")
    instruccion.text = "PAGE"
    fin = OxmlElement("w:fldChar")
    fin.set(qn("w:fldCharType"), "end")
    for elemento in (inicio, instruccion, fin):
        campo._r.append(elemento)


# ---------------------------------------------------------------------------
# Contenido del anexo
# ---------------------------------------------------------------------------

def encabezado(doc: Document) -> None:
    doc.add_heading("Anexo A. Fragmentos de código", level=1)
    parrafo(
        doc,
        "Este anexo recoge los cuatro fragmentos de código que concentran las "
        "decisiones algorítmicas de Polux: el cálculo de los ahorros que guía la "
        "construcción de las rutas, la inserción de las pausas del Reglamento "
        "(CE) nº 561/2006, el cálculo del techo diario de conducción y la fase "
        "de reinserción.",
    )
    parrafo(
        doc,
        "Los fragmentos son **citas literales** del código fuente, con sus "
        "comentarios y sus docstrings en español tal y como están escritos en el "
        "proyecto. Antes de cada bloque se indica el archivo de origen y las "
        "líneas que ocupa en él. La numeración que acompaña al código, en "
        "cambio, **es relativa a cada fragmento** y empieza en 1, para poder "
        "referirse a una línea concreta desde el cuerpo de la memoria.",
    )
    parrafo(
        doc,
        "Cuando un fragmento necesita una constante definida en otro punto del "
        "archivo para entenderse, esa constante se reproduce en un bloque previo "
        "y se indica de dónde procede.",
    )


def a1_ahorros(doc: Document) -> None:
    doc.add_heading("A.1. Cálculo de los ahorros", level=2)
    parrafo(
        doc,
        "El algoritmo de ahorros de Clarke-Wright parte de la solución trivial en "
        "la que cada cliente se atiende con un viaje de ida y vuelta al depósito, "
        "y va fusionando rutas empezando por las fusiones que más kilómetros "
        "ahorran. Esta función calcula ese ahorro para **todos los pares de "
        "clientes** y devuelve la lista ordenada de mayor a menor, que es el "
        "orden en que después se intentan las fusiones.",
    )
    parrafo(
        doc,
        "La expresión implementada en la línea 22 del fragmento es la del "
        "algoritmo original: unir los clientes i y j en una misma ruta evita "
        "regresar al depósito desde i y volver a salir hacia j, a cambio de "
        "recorrer el tramo directo entre ambos.",
    )
    parrafo(
        doc,
        "El detalle que distingue esta implementación de la formulación clásica "
        "es que `matriz` no contiene distancias en línea recta, sino distancias "
        "reales por carretera obtenidas de OSRM. El docstring documenta la "
        "aproximación que eso obliga a asumir: una red viaria real no es "
        "simétrica, y aquí el ahorro se evalúa en un único sentido.",
    )
    procedencia(doc, "algoritmo/clarke_wright.py", 184, 209)
    bloque_codigo(doc, extraer("algoritmo/clarke_wright.py", 184, 209))


def a2_pausas(doc: Document) -> None:
    doc.add_heading("A.2. Inserción de las pausas reglamentarias", level=2)
    parrafo(
        doc,
        "Este es el fragmento donde el Reglamento (CE) nº 561/2006 deja de ser un "
        "texto legal y se convierte en una restricción del modelo. La norma exige "
        "una pausa de 45 minutos tras 4 horas y 30 minutos de conducción "
        "continua, y el punto delicado es que **ese límite puede alcanzarse en "
        "mitad de un trayecto**, no solo al llegar a un cliente: en rutas "
        "interurbanas un único desplazamiento puede durar más de 4 h 30.",
    )
    parrafo(
        doc,
        "Por eso el tramo entre dos nodos no se consume de una vez, sino **por "
        "porciones**: en cada vuelta del bucle se calcula cuánto margen de "
        "conducción continua queda, se recorre como mucho esa cantidad y, si el "
        "margen se agota antes de terminar el tramo, se inserta la pausa y el "
        "contador de conducción continua vuelve a cero. Una versión anterior "
        "comprobaba el límite solo entre tramos, y permitía trayectos de varias "
        "horas con una única pausa.",
    )
    parrafo(
        doc,
        "Cada pausa se registra como un objeto `PausaReglamentaria` con el "
        "instante en que empieza y el tramo en el que cae, que es lo que después "
        "permite imprimirla en su sitio dentro del itinerario de la hoja de ruta "
        "del conductor.",
    )

    parrafo(
        doc,
        "Las constantes que aparecen en el fragmento están definidas al principio "
        "del mismo archivo, en el bloque de constantes del Reglamento:",
    )
    procedencia(
        doc, "algoritmo/restricciones.py", 39, 43,
        nota="constantes citadas por el fragmento",
    )
    bloque_codigo(doc, extraer("algoritmo/restricciones.py", 39, 43))

    parrafo(
        doc,
        "A ellas se añade la tolerancia con la que se comparan los acumulados de "
        "tiempo, necesaria porque son valores en coma flotante y una diferencia "
        "residual bastaría para insertar una pausa espuria o para no terminar "
        "nunca de consumir un tramo:",
    )
    procedencia(
        doc, "algoritmo/restricciones.py", 69, 71,
        nota="constante citada por el fragmento",
    )
    bloque_codigo(doc, extraer("algoritmo/restricciones.py", 69, 71))

    parrafo(
        doc,
        "El fragmento principal es el cuerpo de `simular_ruta`: la inicialización "
        "de los acumuladores, la construcción de la secuencia "
        "depósito → clientes → depósito y el recorrido de cada tramo con la "
        "inserción de pausas. El bucle de porciones ocupa las líneas 32 a 65 del "
        "fragmento.",
    )
    procedencia(doc, "algoritmo/restricciones.py", 170, 235)
    bloque_codigo(doc, extraer("algoritmo/restricciones.py", 170, 235))


def a3_techo(doc: Document) -> None:
    doc.add_heading("A.3. Cálculo del techo diario de conducción", level=2)
    parrafo(
        doc,
        "El Reglamento impone dos límites que actúan a la vez: uno diario "
        "—9 horas, ampliables a 10 como mucho dos veces por semana— y uno "
        "semanal de 56 horas. Un conductor que llegue al viernes con 52 horas "
        "acumuladas no dispone de 9 horas, sino de 4, por mucho que el límite "
        "diario se lo permitiera.",
    )
    parrafo(
        doc,
        "Esta función resuelve esa combinación con el mínimo de ambos límites, y "
        "acota el resultado a cero para que un conductor que ya haya agotado su "
        "semana no reciba un techo negativo. El valor que devuelve es el que "
        "recibe `simular_ruta` como `techo_diario_min`, y es también el que el "
        "resumen previo muestra al usuario antes de optimizar.",
    )
    parrafo(
        doc,
        "Las tres constantes que combina proceden, de nuevo, del bloque de "
        "constantes del Reglamento al principio del archivo:",
    )
    procedencia(
        doc, "algoritmo/restricciones.py", 45, 52,
        nota="constantes citadas por el fragmento",
    )
    bloque_codigo(doc, extraer("algoritmo/restricciones.py", 45, 52))
    procedencia(doc, "algoritmo/restricciones.py", 74, 93)
    bloque_codigo(doc, extraer("algoritmo/restricciones.py", 74, 93))


def a4_reinsercion(doc: Document) -> None:
    doc.add_heading("A.4. Reinserción por inserción más barata", level=2)
    parrafo(
        doc,
        "Tras la mejora local con 2-opt las rutas son más cortas que al terminar "
        "la construcción, de modo que pueden haber liberado holgura suficiente "
        "para atender a algún cliente que se había descartado. La tercera fase "
        "intenta colocar a esos clientes en las rutas ya existentes.",
    )
    parrafo(
        doc,
        "El criterio es el de **inserción más barata**: para cada cliente "
        "pendiente se prueban todas las posiciones de todas las rutas, se valida "
        "cada ruta candidata completa con `simular_ruta` —de forma que la "
        "inserción respeta capacidad, ventanas y Reglamento igual que cualquier "
        "otra ruta— y se conserva la que menos distancia añade.",
    )
    parrafo(
        doc,
        "El bucle exterior repite el barrido mientras se consiga colocar algún "
        "cliente, porque cada inserción cambia la holgura disponible para las "
        "siguientes: un cliente que no cabía en la primera pasada puede caber "
        "después de que otro haya entrado en una ruta distinta.",
    )
    procedencia(
        doc, "algoritmo/reinsercion.py", 86, 109,
        nota="bucle que repite mientras haya inserciones",
    )
    bloque_codigo(doc, extraer("algoritmo/reinsercion.py", 86, 109))

    parrafo(
        doc,
        "La búsqueda propiamente dicha recorre cada ruta y cada posición, "
        "comparando el incremento de distancia que provoca cada inserción "
        "factible:",
    )
    procedencia(doc, "algoritmo/reinsercion.py", 119, 168)
    bloque_codigo(doc, extraer("algoritmo/reinsercion.py", 119, 168))

    parrafo(
        doc,
        "Conviene señalar una limitación deliberada, documentada también en la "
        "memoria: esta fase **solo inserta en rutas que ya existen**. Nunca abre "
        "una ruta nueva aunque quede algún vehículo libre, ni intercambia "
        "clientes entre rutas.",
    )


def main() -> int:
    doc = Document()
    corregir_ajustes(doc)
    preparar_estilos(doc)
    pie_de_pagina(doc)

    encabezado(doc)
    a1_ahorros(doc)
    a2_pausas(doc)
    a3_techo(doc)
    a4_reinsercion(doc)

    doc.save(DESTINO)
    print(f"Anexo generado: {DESTINO}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
