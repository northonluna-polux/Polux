"""
Genera el anexo B de la memoria (tablas de resultados) en formato Word.

Las tablas se construyen leyendo los archivos de resultados que ya
existen en el repositorio; este script no ejecuta ningún experimento ni
recalcula ninguna métrica medida:

    resultados/benchmark_solomon.csv        -> tabla B.1
    resultados/benchmark_flota_acotada.csv  -> tabla B.2
    datos/benchmark/mejores_conocidos.csv   -> tabla B.3 (y referencias B.1)

Lo único que se deriva aquí son magnitudes que se obtienen por aritmética
directa de columnas ya presentes: la desviación porcentual en distancia
frente a la mejor solución conocida, la diferencia de clientes atendidos
entre las dos variantes y las filas de resumen por clase.

Uso:
    python scripts/generar_anexo_resultados.py

Genera `docs/ANEXO_B_RESULTADOS.docx`. Requiere `python-docx`, que solo se
necesita para producir documentación, no para ejecutar Polux.
"""

import csv
import os
from collections import defaultdict

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DESTINO = os.path.join(RAIZ, "docs", "ANEXO_B_RESULTADOS.docx")

RUTA_ILIMITADA = "resultados/benchmark_solomon.csv"
RUTA_ACOTADA = "resultados/benchmark_flota_acotada.csv"
RUTA_MEJORES = "datos/benchmark/mejores_conocidos.csv"

AZUL_POLUX = RGBColor(0x1F, 0x3A, 0x68)
AZUL_CABECERA = "1F3A68"
GRIS_RESUMEN = "E4E9F0"

#: Orden en que se presentan las clases de Solomon, que es el habitual en
#: la literatura: agrupadas (C), aleatorias (R) y mixtas (RC), y dentro de
#: cada una la serie 1 (ventanas estrechas) antes que la serie 2.
ORDEN_CLASES = ["C1", "C2", "R1", "R2", "RC1", "RC2"]

TAMANO_TABLA_PT = 8

#: Medidas de una hoja A4. La plantilla de python-docx trae tamaño carta,
#: así que hay que fijarlas de forma explícita en cada sección.
A4_CORTO = Cm(21.0)
A4_LARGO = Cm(29.7)


# ---------------------------------------------------------------------------
# Formato numérico español
# ---------------------------------------------------------------------------

def num(valor: float, decimales: int = 2) -> str:
    """
    Formatea un número con la convención española: coma decimal y punto
    como separador de millares.

    Se formatea primero con la convención inglesa y se intercambian los
    separadores usando un carácter intermedio, porque hacerlo con dos
    reemplazos encadenados convertiría las comas recién puestas.
    """
    texto = f"{valor:,.{decimales}f}"
    return texto.replace(",", "\x00").replace(".", ",").replace("\x00", ".")


def porcentaje(valor: float, decimales: int = 1) -> str:
    """Formatea una desviación porcentual con su signo explícito."""
    return f"{'+' if valor >= 0 else '−'}{num(abs(valor), decimales)}"


# ---------------------------------------------------------------------------
# Lectura de los archivos de resultados
# ---------------------------------------------------------------------------

def leer_csv(ruta_relativa: str) -> list[dict]:
    ruta = os.path.join(RAIZ, ruta_relativa)
    if not os.path.isfile(ruta):
        raise SystemExit(f"No se encuentra {ruta_relativa}.")
    with open(ruta, encoding="utf-8") as archivo:
        return list(csv.DictReader(archivo))


def indexar_por_instancia(filas: list[dict]) -> dict[str, dict[str, dict]]:
    """Agrupa las filas por instancia y por variante (con/sin reinserción)."""
    indice: dict[str, dict[str, dict]] = defaultdict(dict)
    for fila in filas:
        indice[fila["instancia"]][fila["reinsercion"]] = fila
    return indice


def leer_mejores_conocidos() -> dict[str, dict]:
    """
    Indexa las mejores soluciones publicadas por nombre de instancia.

    La clave se normaliza a minúsculas: este archivo nombra las instancias
    en minúscula (`c101`) y los archivos de resultados en mayúscula
    (`C101`).
    """
    return {fila["instancia"].lower(): fila for fila in leer_csv(RUTA_MEJORES)}


def ordenar(instancias) -> list[str]:
    """Ordena por clase y, dentro de cada clase, por nombre de instancia."""
    return sorted(instancias, key=lambda nombre: (ORDEN_CLASES.index(_clase(nombre)), nombre))


def _clase(nombre: str) -> str:
    """Deduce la clase (C1, RC2, ...) del nombre de la instancia."""
    letras = "".join(caracter for caracter in nombre if caracter.isalpha()).upper()
    serie = "".join(caracter for caracter in nombre if caracter.isdigit())[0]
    return f"{letras}{serie}"


# ---------------------------------------------------------------------------
# Utilidades de maquetación
# ---------------------------------------------------------------------------

def corregir_ajustes(documento: Document) -> None:
    """Completa el `w:zoom` que la plantilla de python-docx deja incompleto."""
    zoom = documento.settings.element.find(qn("w:zoom"))
    if zoom is not None and zoom.get(qn("w:percent")) is None:
        zoom.set(qn("w:percent"), "100")


def preparar_estilos(documento: Document) -> None:
    normal = documento.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.15

    for nivel, tamano in ((1, 18), (2, 14)):
        estilo = documento.styles[f"Heading {nivel}"]
        estilo.font.name = "Calibri"
        estilo.font.size = Pt(tamano)
        estilo.font.color.rgb = AZUL_POLUX
        estilo.font.bold = True
        estilo.paragraph_format.space_before = Pt(18 if nivel == 1 else 14)
        estilo.paragraph_format.space_after = Pt(6)


def apaisar(seccion) -> None:
    """
    Pone la sección en A4 horizontal.

    Se fijan las medidas de forma explícita porque la plantilla por
    defecto de python-docx usa el tamaño carta, y hay que asignar ancho y
    alto además de la orientación: el atributo por sí solo no redimensiona
    la página.
    """
    seccion.orientation = WD_ORIENT.LANDSCAPE
    seccion.page_width = A4_LARGO
    seccion.page_height = A4_CORTO
    seccion.top_margin = Cm(1.8)
    seccion.bottom_margin = Cm(1.8)
    seccion.left_margin = Cm(1.6)
    seccion.right_margin = Cm(1.6)


def enderezar(seccion) -> None:
    """Pone la sección en A4 vertical, para las tablas estrechas."""
    seccion.orientation = WD_ORIENT.PORTRAIT
    seccion.page_width = A4_CORTO
    seccion.page_height = A4_LARGO
    seccion.top_margin = Cm(2.2)
    seccion.bottom_margin = Cm(2.2)
    seccion.left_margin = Cm(2.0)
    seccion.right_margin = Cm(2.0)


def _trocear(texto: str) -> list[dict]:
    trozos, actual, negrita, codigo, indice = [], "", False, False, 0
    while indice < len(texto):
        if texto.startswith("**", indice):
            if actual:
                trozos.append({"texto": actual, "negrita": negrita, "codigo": codigo})
                actual = ""
            negrita = not negrita
            indice += 2
        elif texto[indice] == "`":
            if actual:
                trozos.append({"texto": actual, "negrita": negrita, "codigo": codigo})
                actual = ""
            codigo = not codigo
            indice += 1
        else:
            actual += texto[indice]
            indice += 1
    if actual:
        trozos.append({"texto": actual, "negrita": negrita, "codigo": codigo})
    return trozos


def parrafo(documento: Document, texto: str, tamano: int = 11):
    p = documento.add_paragraph()
    for trozo in _trocear(texto):
        run = p.add_run(trozo["texto"])
        run.bold = trozo["negrita"]
        run.font.size = Pt(tamano)
        if trozo["codigo"]:
            run.font.name = "Consolas"
            run.font.size = Pt(tamano - 1)
    return p


#: Hijos de `w:tcPr` que el esquema de OOXML coloca después de `w:shd`.
#: Sombrear una celda a la que ya se le ha fijado la alineación vertical
#: dejaría el `w:shd` detrás del `w:vAlign`, en un orden que el esquema no
#: admite y que algunos procesadores resuelven ignorando el color.
POSTERIORES_A_SHD = ("noWrap", "tcMar", "textDirection", "tcFitText", "vAlign", "hideMark")


def _sombrear(celda, color: str) -> None:
    sombreado = OxmlElement("w:shd")
    sombreado.set(qn("w:val"), "clear")
    sombreado.set(qn("w:fill"), color)

    propiedades = celda._tc.get_or_add_tcPr()
    posteriores = {qn(f"w:{etiqueta}") for etiqueta in POSTERIORES_A_SHD}
    for posicion, hijo in enumerate(propiedades):
        if hijo.tag in posteriores:
            propiedades.insert(posicion, sombreado)
            return
    propiedades.append(sombreado)


def _repetir_cabecera(fila) -> None:
    """Marca la fila como cabecera para que se repita en cada página."""
    propiedades = fila._tr.get_or_add_trPr()
    cabecera = OxmlElement("w:tblHeader")
    cabecera.set(qn("w:val"), "true")
    propiedades.append(cabecera)


def _no_partir(fila) -> None:
    """Impide que una fila se reparta entre dos páginas."""
    propiedades = fila._tr.get_or_add_trPr()
    entera = OxmlElement("w:cantSplit")
    propiedades.append(entera)


def _pegar_a_la_siguiente(fila) -> None:
    """
    Pide que la fila se mantenga en la misma página que la siguiente.

    Aplicado a todas las filas de un bloque de clase menos a la última, el
    bloque entero viaja junto y no queda partido entre dos páginas salvo
    que no quepa en ninguna.
    """
    for celda in fila.cells:
        for p in celda.paragraphs:
            p.paragraph_format.keep_with_next = True


def escribir_celda(celda, texto: str, negrita: bool = False,
                   alineacion=WD_ALIGN_PARAGRAPH.RIGHT) -> None:
    celda.text = ""
    celda.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = celda.paragraphs[0]
    p.alignment = alineacion
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    run = p.add_run(texto)
    run.bold = negrita
    run.font.size = Pt(TAMANO_TABLA_PT)


def construir_tabla(documento: Document, cabeceras: list[str],
                    bloques: list[dict], anchos_cm: list[float],
                    columnas_texto: set[int]):
    """
    Construye una tabla con la cabecera repetida y un bloque por clase.

    Cada bloque es {"filas": [...], "resumen": [...]}: las filas de sus
    instancias y su fila de resumen, que se sombrea para distinguirla.

    `columnas_texto` indica qué columnas llevan texto y se alinean a la
    izquierda; las demás son numéricas y se alinean a la derecha, que es
    como se comparan las cifras de un vistazo.
    """
    tabla = documento.add_table(rows=1, cols=len(cabeceras))
    tabla.style = "Table Grid"
    tabla.autofit = False

    fila_cabecera = tabla.rows[0]
    _repetir_cabecera(fila_cabecera)
    _no_partir(fila_cabecera)
    for celda, titulo in zip(fila_cabecera.cells, cabeceras):
        escribir_celda(celda, titulo, negrita=True, alineacion=WD_ALIGN_PARAGRAPH.CENTER)
        _sombrear(celda, AZUL_CABECERA)
        for p in celda.paragraphs:
            for run in p.runs:
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    for bloque in bloques:
        filas_bloque = []
        for valores in bloque["filas"]:
            fila = tabla.add_row()
            _no_partir(fila)
            filas_bloque.append(fila)
            for indice, (celda, valor) in enumerate(zip(fila.cells, valores)):
                alineacion = (WD_ALIGN_PARAGRAPH.LEFT if indice in columnas_texto
                              else WD_ALIGN_PARAGRAPH.RIGHT)
                escribir_celda(celda, valor, alineacion=alineacion)

        fila_resumen = tabla.add_row()
        _no_partir(fila_resumen)
        filas_bloque.append(fila_resumen)
        for indice, (celda, valor) in enumerate(zip(fila_resumen.cells, bloque["resumen"])):
            alineacion = (WD_ALIGN_PARAGRAPH.LEFT if indice in columnas_texto
                          else WD_ALIGN_PARAGRAPH.RIGHT)
            escribir_celda(celda, valor, negrita=True, alineacion=alineacion)
            _sombrear(celda, GRIS_RESUMEN)

        # Todas las filas del bloque menos la última se pegan a la
        # siguiente, de modo que la clase no se parta entre dos páginas.
        for fila in filas_bloque[:-1]:
            _pegar_a_la_siguiente(fila)

    for fila in tabla.rows:
        for celda, ancho in zip(fila.cells, anchos_cm):
            celda.width = Cm(ancho)

    documento.add_paragraph().paragraph_format.space_after = Pt(2)
    return tabla


def nota_tabla(documento: Document, texto: str) -> None:
    p = documento.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    for trozo in _trocear(texto):
        run = p.add_run(trozo["texto"])
        run.bold = trozo["negrita"]
        run.italic = not trozo["negrita"]
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        if trozo["codigo"]:
            run.font.name = "Consolas"
            run.font.size = Pt(7.5)


def pie_de_pagina(seccion) -> None:
    p = seccion.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Anexo B — Resultados por instancia   ·   ")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(0x77, 0x77, 0x77)

    campo = p.add_run()
    campo.font.size = Pt(8)
    campo.font.color.rgb = RGBColor(0x77, 0x77, 0x77)
    inicio = OxmlElement("w:fldChar")
    inicio.set(qn("w:fldCharType"), "begin")
    instruccion = OxmlElement("w:instrText")
    instruccion.set(qn("xml:space"), "preserve")
    instruccion.text = "PAGE"
    fin = OxmlElement("w:fldChar")
    fin.set(qn("w:fldCharType"), "end")
    for elemento in (inicio, instruccion, fin):
        campo._r.append(elemento)


# ---------------------------------------------------------------------------
# Tablas
# ---------------------------------------------------------------------------

def tabla_b1(doc: Document) -> None:
    doc.add_heading("B.1. Flota ilimitada: resultados por instancia", level=2)
    parrafo(
        doc,
        "Resultados de las 56 instancias de Solomon de 100 clientes ejecutadas en "
        "**modo de flota ilimitada**, es decir, permitiendo tantos vehículos como "
        "haga falta para atender a todos los clientes. En esta configuración las "
        "56 instancias atienden a los 100 clientes, de modo que las columnas de "
        "clientes servidos y no asignados no aportan información y se omiten.",
    )
    parrafo(
        doc,
        "Los tiempos están en minutos y las distancias en las unidades de las "
        "instancias de Solomon, en las que el tiempo de viaje coincide "
        "numéricamente con la distancia euclídea.",
    )

    ilimitada = indexar_por_instancia(leer_csv(RUTA_ILIMITADA))
    mejores = leer_mejores_conocidos()

    bloques, acumulado = [], defaultdict(list)
    for clase in ORDEN_CLASES:
        instancias = [i for i in ordenar(ilimitada) if ilimitada[i]["sí"]["clase"] == clase]
        filas, columnas = [], defaultdict(list)

        for nombre in instancias:
            fila = ilimitada[nombre]["sí"]
            referencia = mejores[nombre.lower()]

            vehiculos = int(fila["num_vehiculos"])
            programacion = float(fila["tiempo_total_programacion"])
            distancia = float(fila["distancia_total"])
            espera = float(fila["tiempo_espera_total"])
            computo = float(fila["tiempo_computo_segundos"])
            vehiculos_ref = int(referencia["num_vehiculos"])
            distancia_ref = float(referencia["distancia_total"])
            desviacion = (distancia - distancia_ref) / distancia_ref * 100

            for clave, valor in (
                ("vehiculos", vehiculos), ("programacion", programacion),
                ("distancia", distancia), ("espera", espera), ("computo", computo),
                ("vehiculos_ref", vehiculos_ref), ("distancia_ref", distancia_ref),
                ("desviacion", desviacion),
            ):
                columnas[clave].append(valor)
                acumulado[clave].append(valor)

            filas.append([
                nombre, clase, num(vehiculos, 0), num(programacion), num(distancia),
                num(espera), num(computo, 4), num(vehiculos_ref, 0),
                num(distancia_ref), porcentaje(desviacion),
            ])

        def media(clave: str) -> float:
            return sum(columnas[clave]) / len(columnas[clave])

        bloques.append({
            "filas": filas,
            "resumen": [
                f"Media {clase}", f"{len(filas)} inst.",
                num(media("vehiculos"), 1), num(media("programacion")),
                num(media("distancia")), num(media("espera")),
                num(media("computo"), 4), num(media("vehiculos_ref"), 1),
                num(media("distancia_ref")), porcentaje(media("desviacion")),
            ],
        })

    construir_tabla(
        doc,
        ["Instancia", "Clase", "Vehículos", "Tiempo total de\nprogramación (min)",
         "Distancia total", "Tiempo de\nespera (min)", "Tiempo de\ncómputo (s)",
         "Vehículos de\nreferencia", "Distancia de\nreferencia", "Desviación en\ndistancia (%)"],
        bloques,
        [2.2, 1.3, 1.9, 3.0, 2.4, 2.4, 2.4, 2.4, 2.6, 2.6],
        columnas_texto={0, 1},
    )
    nota_tabla(
        doc,
        "Las columnas de referencia proceden de las mejores soluciones publicadas "
        "(tabla B.3), no del archivo de resultados. La desviación en distancia es "
        "(distancia obtenida − distancia de referencia) / distancia de referencia, "
        "y la de las filas de resumen es la media de las desviaciones de las "
        "instancias de la clase, no la desviación de las medias.",
    )
    media_global = sum(acumulado["desviacion"]) / len(acumulado["desviacion"])
    nota_tabla(
        doc,
        f"Media de las 56 instancias: {num(sum(acumulado['vehiculos']) / 56, 1)} vehículos, "
        f"{num(sum(acumulado['distancia']) / 56)} de distancia frente a "
        f"{num(sum(acumulado['distancia_ref']) / 56)} de referencia, "
        f"desviación media {porcentaje(media_global)} %.",
    )


def tabla_b2(doc: Document) -> None:
    doc.add_heading("B.2. Flota acotada: aportación de la fase de reinserción", level=2)
    parrafo(
        doc,
        "Las mismas 56 instancias ejecutadas con la **flota acotada** al número de "
        "vehículos de la mejor solución publicada. Con la flota limitada ya no se "
        "puede atender a todos los clientes, y es entonces cuando la fase de "
        "reinserción tiene margen para aportar algo: la tabla compara los clientes "
        "atendidos desactivándola y activándola.",
    )

    acotada = indexar_por_instancia(leer_csv(RUTA_ACOTADA))
    bloques, total = [], defaultdict(int)

    for clase in ORDEN_CLASES:
        instancias = [i for i in ordenar(acotada) if acotada[i]["sí"]["clase"] == clase]
        filas, suma = [], defaultdict(int)

        for nombre in instancias:
            con = acotada[nombre]["sí"]
            sin = acotada[nombre]["no"]

            permitidos = int(con["num_vehiculos_permitidos"])
            servidos_sin = int(sin["clientes_servidos"])
            servidos_con = int(con["clientes_servidos"])
            diferencia = servidos_con - servidos_sin
            variacion = diferencia / servidos_sin * 100 if servidos_sin else 0.0
            flota = int(con["sin_asignar_flota"])

            for clave, valor in (
                ("permitidos", permitidos), ("sin", servidos_sin),
                ("con", servidos_con), ("diferencia", diferencia), ("flota", flota),
            ):
                suma[clave] += valor
                total[clave] += valor

            filas.append([
                nombre, clase, num(permitidos, 0), num(servidos_sin, 0),
                num(servidos_con, 0), num(diferencia, 0), porcentaje(variacion),
                num(flota, 0),
            ])

        variacion_clase = suma["diferencia"] / suma["sin"] * 100 if suma["sin"] else 0.0
        bloques.append({
            "filas": filas,
            "resumen": [
                f"Total {clase}", f"{len(filas)} inst.",
                num(suma["permitidos"], 0), num(suma["sin"], 0), num(suma["con"], 0),
                num(suma["diferencia"], 0), porcentaje(variacion_clase),
                num(suma["flota"], 0),
            ],
        })

    construir_tabla(
        doc,
        ["Instancia", "Clase", "Vehículos\npermitidos", "Clientes atendidos\nsin reinserción",
         "Clientes atendidos\ncon reinserción", "Diferencia", "Diferencia (%)",
         "No asignados por\nmotivo FLOTA"],
        bloques,
        [2.4, 1.6, 2.6, 3.6, 3.6, 2.4, 2.6, 3.4],
        columnas_texto={0, 1},
    )
    nota_tabla(
        doc,
        "Las filas de resumen son **totales** de la clase, no medias, porque lo que "
        "interesa de esta tabla es cuántos clientes se recuperan en conjunto. La "
        "diferencia porcentual de esas filas se calcula sobre los totales.",
    )
    nota_tabla(
        doc,
        "La columna de no asignados por motivo FLOTA corresponde a la ejecución "
        "**con** reinserción, que es la configuración por defecto de la aplicación.",
    )
    variacion_global = total["diferencia"] / total["sin"] * 100
    nota_tabla(
        doc,
        f"En conjunto: {num(total['sin'], 0)} clientes atendidos sin reinserción y "
        f"{num(total['con'], 0)} con ella, {num(total['diferencia'], 0)} más "
        f"({porcentaje(variacion_global)} %).",
    )


def tabla_b3(doc: Document) -> None:
    doc.add_heading("B.3. Mejores soluciones conocidas usadas como referencia", level=2)
    parrafo(
        doc,
        "Mejores soluciones publicadas para las 56 instancias, tal y como las "
        "recoge SINTEF, que es el repositorio de referencia del VRPTW. Son los "
        "valores con los que se comparan los resultados de la tabla B.1 y los que "
        "fijan el número de vehículos permitidos en la tabla B.2. El criterio de "
        "la literatura es jerárquico: primero se minimiza el número de vehículos "
        "y, a igualdad de vehículos, la distancia.",
    )

    #: Particularidades de etiquetado de tres instancias. No son una columna
    #: del archivo: proceden de cómo publica SINTEF esos tres valores, y se
    #: anotan aquí para que la cifra de la tabla pueda contrastarse con la
    #: fuente sin que parezca un error de transcripción.
    marcas = {"r111": "a", "r205": "b", "rc101": "b"}

    mejores = leer_mejores_conocidos()
    bloques = []
    for clase in ORDEN_CLASES:
        instancias = [i for i in ordenar(mejores) if _clase(i) == clase]
        filas, vehiculos_clase, distancias_clase = [], [], []

        for nombre in instancias:
            fila = mejores[nombre]
            vehiculos = int(fila["num_vehiculos"])
            distancia = float(fila["distancia_total"])
            vehiculos_clase.append(vehiculos)
            distancias_clase.append(distancia)

            marca = marcas.get(nombre, "")
            filas.append([
                nombre + (f" ({marca})" if marca else ""),
                num(vehiculos, 0), num(distancia), fila["fuente"],
            ])

        bloques.append({
            "filas": filas,
            "resumen": [
                f"Media {clase} ({len(filas)} inst.)",
                num(sum(vehiculos_clase) / len(vehiculos_clase), 1),
                num(sum(distancias_clase) / len(distancias_clase)),
                "—",
            ],
        })

    construir_tabla(
        doc,
        ["Instancia", "Vehículos", "Distancia", "Fuente"],
        bloques,
        [4.2, 2.8, 3.2, 6.8],
        columnas_texto={0, 3},
    )
    nota_tabla(
        doc,
        "**(a)** En `r111` la distancia publicada aparece acompañada de una llamada "
        "a pie de página referida al propio valor de la distancia.",
    )
    nota_tabla(
        doc,
        "**(b)** `r205` y `rc101` se publican como `r205b` y `rc101b`: las cifras "
        "comunicadas originalmente se consideraron afectadas por errores de "
        "redondeo, y la variante «b» es la que recoge el valor corregido.",
    )
    nota_tabla(
        doc,
        "Las siglas de la columna «Fuente» son las que emplea SINTEF para "
        "identificar al equipo que publicó cada solución.",
    )


def main() -> int:
    doc = Document()
    corregir_ajustes(doc)
    preparar_estilos(doc)

    apaisar(doc.sections[0])
    pie_de_pagina(doc.sections[0])

    doc.add_heading("Anexo B. Resultados por instancia", level=1)
    parrafo(
        doc,
        "Este anexo recoge los resultados completos de la validación experimental, "
        "instancia a instancia. Las tablas se generan leyendo directamente los "
        "archivos de resultados del repositorio, sin volver a ejecutar los "
        "experimentos.",
    )
    parrafo(
        doc,
        "Todas las cifras siguen la convención española: coma decimal y punto como "
        "separador de millares.",
    )

    tabla_b1(doc)
    doc.add_section(WD_SECTION.NEW_PAGE)
    apaisar(doc.sections[-1])
    tabla_b2(doc)
    # La tabla B.3 solo tiene cuatro columnas: en horizontal quedaría una
    # columna «Fuente» absurdamente ancha, así que va en vertical.
    doc.add_section(WD_SECTION.NEW_PAGE)
    enderezar(doc.sections[-1])
    tabla_b3(doc)

    doc.save(DESTINO)
    print(f"Anexo generado: {DESTINO}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
